"""Exportación a Excel (openpyxl). Reemplaza los 'Exportar a Excel' del VFP."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment

AZUL = "FF14428A"


def _encabezar(ws, columnas):
    ws.append(columnas)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(horizontal="center")


ESTADO_TXT = {"A": "Activo", "C": "Cancelado"}


def _num(ws, min_col, max_col, min_row=2):
    for row in ws.iter_rows(min_col=min_col, max_col=max_col, min_row=min_row):
        for cell in row:
            cell.number_format = "#,##0.00"


def _anchos(ws, anchos):
    for idx, w in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + idx)].width = w


def _guardar(wb) -> bytes:
    buf = BytesIO(); wb.save(buf); buf.seek(0); return buf.read()


def sumas_y_saldos_excel(data: dict) -> bytes:
    """Balance de comprobación: por cuenta con debe/haber y saldo deudor/acreedor + totales."""
    wb = Workbook(); ws = wb.active; ws.title = "Sumas y saldos"
    _encabezar(ws, ["Código", "Cuenta", "Rubro", "Debe", "Haber", "Saldo deudor", "Saldo acreedor"])
    for f in data.get("filas", []):
        ws.append([f["codigo"], f["nombre"], (f.get("tipo") or "").capitalize(),
                   float(f["debe"]), float(f["haber"]), float(f["saldo_deudor"]), float(f["saldo_acreedor"])])
    t = data.get("totales", {})
    ws.append([])
    ws.append(["", "", "TOTALES", float(t.get("debe", 0)), float(t.get("haber", 0)),
               float(t.get("saldo_deudor", 0)), float(t.get("saldo_acreedor", 0))])
    _num(ws, 4, 7)
    _anchos(ws, [12, 36, 13, 15, 15, 15, 15])
    return _guardar(wb)


def estados_contables_excel(data: dict) -> bytes:
    """Situación patrimonial + Estado de resultados en dos hojas."""
    wb = Workbook()
    sit = data.get("situacion", {}); res = data.get("resultados", {})
    ws = wb.active; ws.title = "Situación patrimonial"
    _encabezar(ws, ["Rubro", "Código", "Cuenta", "Importe"])
    def _bloque(hoja, titulo, cuentas, total):
        hoja.append([titulo, "", "", float(total)])
        hoja[hoja.max_row][0].font = Font(bold=True)
        for c in (cuentas or []):
            hoja.append(["", c["codigo"], c["nombre"], float(c["valor"])])
    _bloque(ws, "ACTIVO", sit.get("activo", {}).get("cuentas"), sit.get("activo", {}).get("total", 0))
    ws.append([])
    _bloque(ws, "PASIVO", sit.get("pasivo", {}).get("cuentas"), sit.get("pasivo", {}).get("total", 0))
    pn = sit.get("patrimonio", {})
    _bloque(ws, "PATRIMONIO NETO", (pn.get("cuentas") or []) + [{"codigo": "—", "nombre": "Resultado del ejercicio", "valor": pn.get("resultado_ejercicio", 0)}], pn.get("total", 0))
    if sit.get("otros", {}).get("cuentas"):
        ws.append([])
        _bloque(ws, "OTRAS (a clasificar)", sit["otros"]["cuentas"], sit["otros"].get("total", 0))
    ws.append([])
    ws.append(["Activo", "", "", float(sit.get("total_activo", 0))])
    ws.append(["Pasivo + PN", "", "", float(sit.get("total_pasivo_pn", 0))])
    _num(ws, 4, 4); _anchos(ws, [22, 12, 34, 16])

    wr = wb.create_sheet("Estado de resultados")
    _encabezar(wr, ["Rubro", "Código", "Cuenta", "Importe"])
    _bloque(wr, "INGRESOS", res.get("ingresos", {}).get("cuentas"), res.get("ingresos", {}).get("total", 0))
    wr.append([])
    _bloque(wr, "EGRESOS", res.get("egresos", {}).get("cuentas"), res.get("egresos", {}).get("total", 0))
    wr.append([])
    wr.append(["RESULTADO DEL EJERCICIO", "", "", float(res.get("resultado", 0))])
    wr[wr.max_row][0].font = Font(bold=True)
    _num(wr, 4, 4); _anchos(wr, [26, 12, 34, 16])
    return _guardar(wb)


def libro_diario_excel(asientos: list) -> bytes:
    """Libro diario: una fila por línea de asiento (N°, fecha, diario, concepto, cuenta, debe, haber)."""
    wb = Workbook(); ws = wb.active; ws.title = "Libro diario"
    _encabezar(ws, ["N°", "Fecha", "Diario", "Concepto", "Cuenta", "Debe", "Haber", "Estado"])
    for a in asientos:
        for l in (a.lineas or []):
            ws.append([a.numero or a.id, str(a.fecha), a.diario_codigo or "", a.concepto,
                       f"{l.cuenta_codigo} · {l.cuenta_nombre}", float(l.debe), float(l.haber), a.estado])
    _num(ws, 6, 7)
    _anchos(ws, [7, 12, 10, 34, 34, 15, 15, 11])
    return _guardar(wb)


def listado_creditos_excel(items) -> bytes:
    """Listado de créditos: crédito, cliente, CUIL, línea, capital, saldo, estado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Créditos"
    _encabezar(ws, ["Crédito", "Cliente", "CUIL", "Línea", "Capital", "Saldo", "Estado"])
    tot_cap = tot_sal = 0.0
    for c in items:
        cap, sal = float(c["capital"]), float(c["saldo"])
        tot_cap += cap; tot_sal += sal
        ws.append([c["credito_id"], c["cliente"], c["cuil"], c["linea"],
                   cap, sal, ESTADO_TXT.get(c["estado"], c["estado"])])
    ws.append([])
    ws.append(["", "", "", "TOTAL", tot_cap, tot_sal, ""])
    for row in ws.iter_rows(min_col=5, max_col=6, min_row=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    for idx, wdt in enumerate([9, 34, 13, 30, 15, 15, 11], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def pagos_caja_excel(items) -> bytes:
    """Pagos de créditos en caja (write-only: soporta decenas de miles de filas)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Pagos en caja")
    for idx, wdt in enumerate([13, 9, 7, 34, 9, 8, 14, 15], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt

    hdr = []
    for name in ["Fecha pago", "Crédito", "Cuota", "Cliente", "Recibo",
                 "Vía", "Cajero", "Importe"]:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws.append(hdr)

    total = 0.0
    for p in items:
        imp = float(p["total_pagado"])
        total += imp
        importe = WriteOnlyCell(ws, value=imp)
        importe.number_format = "#,##0.00"
        ws.append([str(p["fecha_pago"] or ""), p["credito_id"], p["cuota"],
                   p["cliente"], p["nro_recibo"], p["via_pago"], p["cajero"], importe])
    tot_cell = WriteOnlyCell(ws, value=total)
    tot_cell.number_format = "#,##0.00"
    ws.append(["", "", "", "", "", "", "TOTAL", tot_cell])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def anexo_resolucion_excel(data, lote) -> bytes:
    """Anexo de resolución: listado de solicitudes del lote con capital y total.
    Reconstruye el reporte rpt330851500anexo_disp_a del VFP."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Anexo {lote}"
    ws.append([f"ANEXO {data.get('nombre','')} — Lote/Resolución N° {lote}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    _encabezar(ws, ["N° Sol.", "Fecha", "CUIL", "Solicitante", "Línea",
                    "Denominación", "Capital"])
    total = 0.0
    for s in data["items"]:
        cap = float(s["montosol"]); total += cap
        ws.append([s["no_solicitud"], str(s["fecha_soli"] or ""), s["cuil"],
                   s["apellido_nombre"], s["linea"], s["denominacion"], cap])
    ws.append([])
    ws.append(["", "", "", "", "", "TOTAL", total])
    for row in ws.iter_rows(min_col=7, max_col=7, min_row=4):
        for c in row:
            c.number_format = "#,##0.00"
    for idx, wdt in enumerate([9, 12, 13, 34, 8, 30, 15], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    buf = BytesIO()
    wb.save(buf); buf.seek(0)
    return buf.read()


def ordenes_pago_excel(ops) -> bytes:
    """Informe de órdenes de pago: N°, fecha, beneficiario, CUIT, tipo, concepto,
    importe, estado, cheque."""
    _EST = {"P": "Pendiente", "G": "Girada", "A": "Anulada"}
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Órdenes de pago")
    for idx, wdt in enumerate([8, 12, 34, 13, 12, 30, 15, 11, 12], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    hdr = []
    for name in ["N°", "Fecha", "Beneficiario", "CUIT", "Tipo", "Concepto",
                 "Importe", "Estado", "Cheque"]:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws.append(hdr)
    total = 0.0
    for o in ops:
        imp = float(o.importe)
        if o.estado != "A":
            total += imp
        importe = WriteOnlyCell(ws, value=imp)
        importe.number_format = "#,##0.00"
        ws.append([o.numero, str(o.fecha), o.beneficiario, o.cuit_beneficiario,
                   o.tipo, o.concepto, importe, _EST.get(o.estado, o.estado),
                   o.cheque_numero or ""])
    tot = WriteOnlyCell(ws, value=total)
    tot.number_format = "#,##0.00"
    ws.append(["", "", "", "", "", "TOTAL (excl. anuladas)", tot, "", ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def turnos_excel(items) -> bytes:
    """Turnos otorgados de crédito (write-only): fecha, período, tipo, N°,
    solicitante, CUIL, sueldo, usado, autorizado."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Turnos")
    for idx, wdt in enumerate([12, 9, 8, 8, 34, 13, 14, 8, 12], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    hdr = []
    for name in ["Fecha", "Período", "Tipo", "N° turno", "Solicitante", "CUIL",
                 "Sueldo", "Usado", "Autorizado"]:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws.append(hdr)
    for t in items:
        sueldo = WriteOnlyCell(ws, value=float(t["sueldo"]))
        sueldo.number_format = "#,##0.00"
        ws.append([str(t["fecha"] or ""), t["periodo"], t["tipo"], t["numero"],
                   t["apellido_nombre"], t["cuil"], sueldo,
                   "Sí" if t["usado"] else "No", "Sí" if t["autorizado"] else "No"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def envios_excel(envios) -> bytes:
    """Padrón de débito por planilla (para el banco): CBU, CUIL, cliente, importe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Débitos"
    _encabezar(ws, ["CBU", "CUIL", "Cliente", "Crédito", "Cuota", "Vencimiento", "Importe"])
    total = 0
    for i in envios["items"]:
        imp = float(i["importe"])
        total += imp
        ws.append([i["cbu"], i["cuil"], i["cliente"], i["credito_id"],
                   i["cuota_numero"], str(i["vencimiento"]), imp])
    ws.append([])
    ws.append(["", "", "", "", "", "TOTAL", total])
    for row in ws.iter_rows(min_col=7, max_col=7, min_row=2):
        for c in row:
            c.number_format = "#,##0.00"
    widths = [26, 13, 34, 9, 7, 13, 14]
    for idx, wdt in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def contratos_pp_excel(items) -> bytes:
    """Cartera de contratos de Configurar Créditos: nº, cliente, línea, monto, saldo, tasa, estado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cartera"
    _encabezar(ws, ["Contrato", "Cliente", "Línea", "Sistema", "Monto", "Saldo capital",
                    "TNA %", "Plazo", "Estado", "Fecha valor"])
    tot_monto = tot_saldo = 0.0
    for c in items:
        monto, saldo = float(c["monto_original"]), float(c["saldo_capital"])
        tot_monto += monto; tot_saldo += saldo
        snap = c.get("snapshot") or {}
        ws.append([c["numero_contrato"], c["cliente_nombre"], snap.get("codigo", ""), c["sistema"],
                   monto, saldo, float(c["tasa"]), c["plazo"], c["estado"], c["fecha_valor"]])
    ws.append([])
    ws.append(["", "", "", "TOTAL", tot_monto, tot_saldo, "", "", "", ""])
    for row in ws.iter_rows(min_col=5, max_col=6, min_row=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    for idx, wdt in enumerate([16, 30, 12, 11, 15, 15, 9, 8, 12, 12], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
